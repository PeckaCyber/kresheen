import io
from datetime import UTC, datetime

from docx import Document
from fastapi.testclient import TestClient

from kresheen.dashboard.app import create_app
from kresheen.dashboard.auth import hash_password
from kresheen.models.user import User
from kresheen.storage.sql_repo import SQLRepository

PW = "correct-password-1"  # nosec: test fixture


def test_docx_keeps_its_structure_with_the_fast_paragraph_path(tmp_path):
    url = f"sqlite:///{tmp_path}/t.db"
    SQLRepository(url).save_user(User(id="U1", email="o@x.test", password_hash=hash_password(PW), role="owner", created_at=datetime.now(UTC)))
    c = TestClient(create_app(db_url=url))
    c.headers["X-CSRF-Token"] = c.post("/api/auth/login", json={"email": "o@x.test", "password": PW}).json()["csrf_token"]
    ids = c.post("/api/applications/sample").json()
    r = c.get(f"/api/applications/{ids['application_id']}/runs/{ids['run_id']}/report.docx")
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))  # reopening proves the XML is valid
    styles = [p.style.name for p in doc.paragraphs]
    assert styles[0] == "Title" and doc.paragraphs[0].text == "Sample: online shop"
    assert {"Heading 1", "Heading 2", "List Bullet", "Normal"} <= set(styles)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert "Summary" in headings and any(h.startswith("Threats (") for h in headings)
    assert doc.element.body[-1].tag.endswith("sectPr"), "page settings must stay the last element of the body"
    # order is preserved: Summary comes before the threat list
    texts = [p.text for p in doc.paragraphs]
    assert texts.index("Summary") < next(i for i, t in enumerate(texts) if t.startswith("Threats ("))
    assert len(doc.tables) >= 5
