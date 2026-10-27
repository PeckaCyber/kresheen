"""Report and export endpoints (HTML, DOCX, Markdown, CSV, SVG)."""

import csv
import io
import re
import xml.etree.ElementTree as ET

import pytest
from docx import Document
from fastapi.testclient import TestClient

from kresheen.dashboard.app import create_app
from tests.integration.test_api_applications import _logged_in_client, _run_with_manifest
from tests.integration.test_api_manual_modeling import _add, _threats
from tests.unit.test_markdown_report import _outside_fences

HOSTILE = "<img src=x onerror=alert(1)>"
ENDPOINTS = ("report.html", "report.docx", "report.md", "threats.csv", "findings.csv", "dfd.svg")


def _blank_model_named(client, csrf, name):
    headers = {"X-CSRF-Token": csrf}
    application = client.post("/api/applications", json={"name": name, "description": "d"}, headers=headers).json()
    run = client.post(f"/api/applications/{application['id']}/runs", headers=headers).json()
    return (
        application,
        run,
        f"/api/applications/{application['id']}/runs/{run['id']}",
        f"/api/applications/{application['id']}",
        headers,
    )


@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_exports_require_a_session(tmp_path, endpoint):
    client = TestClient(create_app(db_url=f"sqlite:///{tmp_path}/test.db"))

    assert client.get(f"/api/applications/APP-1/runs/RUN-1/{endpoint}").status_code == 401


@pytest.mark.parametrize("endpoint", ENDPOINTS)
def test_unknown_application_and_run_are_404(tmp_path, endpoint):
    client, csrf = _logged_in_client(tmp_path)
    _application, run, run_url, app_url, _ = _blank_model_named(client, csrf, "Shop")

    assert client.get(f"/api/applications/APP-nope/runs/{run['id']}/{endpoint}").status_code == 404
    assert client.get(f"{app_url}/runs/RUN-nope/{endpoint}").status_code == 404
    assert client.get(f"{run_url}/{endpoint}").status_code == 200


def test_a_run_of_another_application_is_not_readable_through_this_one(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _a, run_a, _ru, _au, _h = _blank_model_named(client, csrf, "First")
    app_b, _run_b, _rub, _aub, _h2 = _blank_model_named(client, csrf, "Second")

    response = client.get(f"/api/applications/{app_b['id']}/runs/{run_a['id']}/report.html")

    assert response.status_code == 404


def test_html_report_content_headers_and_inline_versus_download(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _app, _run, run_url, app_url, headers = _blank_model_named(client, csrf, "Acme Shop")
    _add(client, app_url, headers, "Web")
    _add(client, app_url, headers, "Orders DB", "data_store")

    inline = client.get(f"{run_url}/report.html")

    assert inline.status_code == 200 and inline.headers["content-type"].startswith("text/html")
    assert "content-disposition" not in inline.headers  # viewable, and printable to PDF
    assert "script-src" not in inline.headers["content-security-policy"]
    assert inline.headers["content-security-policy"].startswith("default-src 'none'")
    assert inline.headers["cache-control"] == "no-store" and inline.headers["x-content-type-options"] == "nosniff"
    body = inline.text
    assert "Acme Shop" in body and "Orders DB" in body and "<svg" in body
    assert "<script" not in body.lower()
    assert "Print" in body  # the how-to-get-a-PDF hint
    assert "None (a model drawn by hand)" in body
    downloaded = client.get(f"{run_url}/report.html?download=1")
    assert downloaded.headers["content-disposition"].startswith('attachment; filename="Acme-Shop-')
    assert downloaded.headers["content-disposition"].endswith('-threat-model.html"')


def test_hostile_names_are_escaped_in_every_format_and_the_filename_stays_safe(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    name = f'{HOSTILE}"\r\nX-Injected: 1'
    _app, run, run_url, app_url, headers = _blank_model_named(client, csrf, name)
    _add(client, app_url, headers, HOSTILE)
    comp = client.get(f"{run_url}/dfd").json()["nodes"][0]
    client.post(
        f"{run_url}/manual-threats",
        headers=headers,
        json={
            "element_kind": "component",
            "element_id": comp["id"],
            "title": HOSTILE,
            "description": HOSTILE,
            "mitigations": [HOSTILE],
        },
    )
    client.put(
        f"{run_url}/threats/status",
        headers=headers,
        json={
            "threat_key": next(t["key"] for t in _threats(client, app_url, {"id": run["id"]})),
            "status": "accepted",
            "note": HOSTILE,
        },
    )

    html = client.get(f"{run_url}/report.html")
    svg = client.get(f"{run_url}/dfd.svg")
    docx = client.get(f"{run_url}/report.docx")
    markdown = client.get(f"{run_url}/report.md")

    # Verbatim only inside fenced blocks and code spans (both literal in Markdown); escaped everywhere else.
    outside = re.sub(r"(`+).+?\1", "", "\n".join(_outside_fences(markdown.text)))
    assert not re.search(r"(?<!\\)<img", outside) and "\\<img src=x" in outside
    assert "#### T1." in markdown.text and "Treat them as data" in markdown.text
    assert markdown.headers["content-type"].startswith("text/markdown")
    assert markdown.headers["content-security-policy"].startswith("default-src 'none'")
    assert markdown.headers["content-disposition"].endswith('-threat-model.md"')
    assert "<img src=x" not in html.text and "&lt;img src=x" in html.text
    assert "<img" not in svg.text
    ET.fromstring(svg.text.split("?>\n", 1)[1])  # still well formed
    text = "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert HOSTILE in text  # literal text in Word, not markup
    for response in (
        html,
        svg,
        docx,
        markdown,
        client.get(f"{run_url}/threats.csv"),
        client.get(f"{run_url}/report.html?download=1"),
    ):
        disposition = response.headers.get("content-disposition", "")
        assert "\n" not in disposition and "\r" not in disposition and "X-Injected" not in response.headers


def test_docx_report_opens_and_contains_the_model(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _app, _run, run_url, app_url, headers = _blank_model_named(client, csrf, "Acme Shop")
    _add(client, app_url, headers, "Web")

    response = client.get(f"{run_url}/report.docx")

    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert response.headers["content-disposition"].startswith('attachment; filename="Acme-Shop-')
    document = Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert "Acme Shop" in text and "Threats (" in text
    assert "Web" in cells and "process" in cells
    assert any("Web" in p.text for p in document.paragraphs)  # a threat section heading names the element


def test_threats_csv_has_the_columns_rows_and_neutralises_formulas(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _app, run, run_url, app_url, headers = _blank_model_named(client, csrf, "Shop")
    web = _add(client, app_url, headers, "=cmd|' /C calc'!A0")
    client.post(
        f"{run_url}/manual-threats",
        headers=headers,
        json={"element_kind": "component", "element_id": web["id"], "title": "@SUM(1+1)", "severity": "high"},
    )

    response = client.get(f"{run_url}/threats.csv")

    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"].endswith('-threats.csv"')
    assert response.text.startswith("\ufeff")
    rows = list(csv.DictReader(io.StringIO(response.text.lstrip("\ufeff"))))
    assert len(rows) == len(_threats(client, app_url, run)) > 0
    assert {r["element"] for r in rows} == {"'=cmd|' /C calc'!A0"}
    assert any(r["title"] == "'@SUM(1+1)" and r["origin"] == "manual" for r in rows)
    assert all(not v.startswith(("=", "+", "@")) for r in rows for v in r.values())


def test_findings_csv_and_report_carry_scanned_findings_and_severity(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    application, run = _run_with_manifest(client, csrf)
    run_url = f"/api/applications/{application['id']}/runs/{run['id']}"

    findings = list(csv.DictReader(io.StringIO(client.get(f"{run_url}/findings.csv").text.lstrip("\ufeff"))))
    report = client.get(f"{run_url}/report.html").text

    assert findings and {"id", "element", "title", "severity", "status"} <= set(findings[0])
    assert findings[0]["title"] in report
    assert "manifest_upload" in report  # the source is listed by type


def test_report_lists_assumptions_and_out_of_scope_with_reasons(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _app, _run, run_url, app_url, headers = _blank_model_named(client, csrf, "Shop")
    db = _add(client, app_url, headers, "Vault", "data_store")
    web = _add(client, app_url, headers, "Legacy")
    client.put(
        f"{run_url}/element-properties",
        headers=headers,
        json={"element_kind": "component", "element_id": db["id"], "properties": {"is_encrypted": False}},
    )
    client.put(
        f"{run_url}/element-properties",
        headers=headers,
        json={
            "element_kind": "component",
            "element_id": web["id"],
            "properties": {"out_of_scope": True, "out_of_scope_reason": "Owned by another team"},
        },
    )

    html = client.get(f"{run_url}/report.html").text

    assert "Encrypted at rest: No" in html
    assert "Owned by another team" in html and "(out of scope)" in html
    assert "owner states: not encrypted at rest" in html  # and the threat it raised, with why


def test_svg_export_is_a_standalone_document_with_restrictive_headers(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _app, _run, run_url, app_url, headers = _blank_model_named(client, csrf, "Shop")
    _add(client, app_url, headers, "Web")

    response = client.get(f"{run_url}/dfd.svg")

    assert response.headers["content-type"].startswith("image/svg+xml")
    assert response.text.startswith("<?xml")
    assert "default-src 'none'" in response.headers["content-security-policy"]
    assert ">Web<" in response.text
    assert "content-disposition" not in response.headers
    assert "diagram.svg" in client.get(f"{run_url}/dfd.svg?download=1").headers["content-disposition"]


def _docs_only_run(client, csrf):
    headers = {"X-CSRF-Token": csrf}
    application = client.post(
        "/api/applications", json={"name": "Docs only", "description": "d"}, headers=headers
    ).json()
    client.post(
        f"/api/applications/{application['id']}/data-sources",
        json={"type": "documentation", "connection_details": {"filename": "design.md", "content": "# Design"}},
        headers=headers,
    )
    run = client.post(f"/api/applications/{application['id']}/runs", headers=headers).json()
    return application, run


def test_a_documentation_only_run_reports_the_input_as_not_analysed(tmp_path):
    """Stored-but-unread input must never look like a clean result."""
    client, csrf = _logged_in_client(tmp_path)
    application, run = _docs_only_run(client, csrf)

    fetched = client.get(f"/api/applications/{application['id']}/runs/{run['id']}").json()
    assert [r["status"] for r in fetched["source_results"]] == ["reference_only"]

    base = f"/api/applications/{application['id']}/runs/{run['id']}"
    html = client.get(f"{base}/report.html").text
    assert "stored but not analysed" in html and "nothing in this file has been analysed" in html
    assert client.get(f"{base}/report.docx").status_code == 200
    assert client.get(f"{base}/report.pdf").status_code == 200


def test_a_manifest_run_has_no_not_analysed_notice(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    application, run = _run_with_manifest(client, csrf)

    fetched = client.get(f"/api/applications/{application['id']}/runs/{run['id']}").json()
    assert [r["status"] for r in fetched["source_results"]] == ["analysed"]
    html = client.get(f"/api/applications/{application['id']}/runs/{run['id']}/report.html").text
    assert "stored but not analysed" not in html


def test_csv_exports_keep_the_how_to_verify_column(tmp_path):
    client, csrf = _logged_in_client(tmp_path)
    _application, run = _run_with_manifest(client, csrf)
    base = f"/api/applications/{_application['id']}/runs/{run['id']}"

    header = client.get(f"{base}/threats.csv").text.lstrip("\ufeff").splitlines()[0]
    assert "how_to_verify" in header.split(",")
    assert "how_to_verify" in client.get(f"{base}/action-plan.csv").text.splitlines()[0]
