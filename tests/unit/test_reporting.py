"""The report builder and its renderers (reporting/)."""

import csv
import io
import xml.etree.ElementTree as ET
from datetime import UTC, datetime

import pytest
from docx import Document

from kresheen.models.application import Application
from kresheen.models.data_source import DataSourceConfig
from kresheen.models.threat_model_run import ThreatModelRun
from kresheen.reporting.csv_export import THREAT_COLUMNS, safe_cell, threats_csv
from kresheen.reporting.dfd_svg import (
    GROUP_COLUMNS,
    NODE_CELL_H,
    NODE_W,
    layout,
    render_svg,
)
from kresheen.reporting.docx_report import render_docx
from kresheen.reporting.html_report import render_html
from kresheen.reporting.model import (
    ReportBoundary,
    ReportComponent,
    ReportFlow,
    build_report,
    dfd_parts,
    safe_repo_url,
    sbom_vulnerability_rows,
    statements_for,
)
from kresheen.reporting.pdf_report import render_pdf
from kresheen.reporting.sarif_export import build_sarif
from kresheen.reporting.text import clean_text
from kresheen.reporting.threat_dragon_export import build_threat_dragon_model
from kresheen.storage.sql_repo import SQLRepository

HOSTILE = '<img src=x onerror=alert(1)> & "quoted"'


def _component(id_, name, *, type_="process", boundary=None, manual=False, x=None, y=None, **kw):
    return ReportComponent(
        id=id_, name=name, type=type_, trust_boundary=boundary, worst_severity=kw.pop("worst_severity", None),
        open_finding_count=kw.pop("open_finding_count", 0), is_manual=manual, position_x=x, position_y=y, **kw,
    )


def _flow(source, target, label=None, **kw):
    return ReportFlow(
        id=f"F-{source.id}-{target.id}", source_id=source.id, target_id=target.id, source_name=source.name,
        target_name=target.name, label=label, is_manual=kw.pop("is_manual", False), **kw,
    )


# -- diagram ------------------------------------------------------------------


def test_layout_stacks_boundaries_and_grids_members_like_the_screen():
    members = [_component(f"C{i}", f"svc{i}", boundary="prod") for i in range(4)]
    other = _component("O", "vendor", type_="external_entity", boundary="third-party")

    nodes, groups = layout([*members, other], [])

    assert [name for name, _ in groups] == ["prod", "third-party"]
    prod, third = groups[0][1], groups[1][1]
    assert third.y > prod.y + prod.h  # stacked below, not overlapping
    # 4 members in 3 columns: a second row starts at the first column
    assert nodes["C3"].x == nodes["C0"].x and nodes["C3"].y == nodes["C0"].y + NODE_CELL_H
    assert GROUP_COLUMNS == 3
    assert all(prod.x <= nodes[m.id].x and nodes[m.id].x + NODE_W <= prod.x + prod.w for m in members)


def test_layout_honours_saved_boundary_origin_and_hand_placed_positions():
    a = _component("A", "in-zone", boundary="zone")
    b = _component("B", "dropped", manual=True, x=500.0, y=-40.0)
    c = _component("C", "auto")

    nodes, groups = layout([a, b, c], [ReportBoundary(name="zone", position_x=300.0, position_y=200.0)])

    assert (groups[0][1].x, groups[0][1].y) == (300.0, 200.0)
    assert (nodes["B"].x, nodes["B"].y) == (500.0, -40.0)
    assert nodes["C"].y > groups[0][1].y  # gridded below the boundaries


def test_svg_is_well_formed_and_escapes_every_user_string():
    web = _component("W", HOSTILE, boundary=HOSTILE)
    db = _component("D", "DB", type_="data_store")

    svg = render_svg([web, db], [_flow(web, db, label=HOSTILE)], [])

    root = ET.fromstring(svg.split("?>\n", 1)[1])  # parses: not broken by the quotes and angle brackets
    assert root.tag.endswith("svg")
    assert "<img" not in svg and "onerror=alert(1)>" not in svg
    assert "&lt;img" in svg
    texts = " ".join(t.text or "" for t in root.iter() if t.tag.endswith("text"))
    assert "<img src=x" in texts  # the literal text survives, as text


def test_svg_marks_out_of_scope_and_severity_and_edges_use_node_borders():
    a = _component("A", "api", worst_severity="critical", open_finding_count=3)
    b = _component("B", "legacy", out_of_scope=True)

    svg = render_svg([a, b], [_flow(a, b, label="HTTP")], [])

    assert "#dc3545" in svg  # critical fill
    assert "api (3)" in svg
    assert "out of scope" in svg and "opacity=\"0.45\"" in svg
    assert "marker-end" in svg and ">HTTP<" in svg


def test_svg_of_an_empty_model_says_so_and_is_valid():
    svg = render_svg([], [], [], standalone=False)

    assert not svg.startswith("<?xml")
    assert ET.fromstring(svg) is not None
    assert "No components in this model." in svg


def test_svg_skips_flows_whose_endpoint_is_missing_and_wraps_long_names():
    a = _component("A", "x" * 80)

    svg = render_svg([a], [ReportFlow(id="F", source_id="A", target_id="GONE", source_name="x", target_name="?",
                                      label=None, is_manual=False)], [])

    assert "marker-end" not in svg
    root = ET.fromstring(svg.split("?>\n", 1)[1])
    lines = [t.text for t in root.iter() if t.tag.endswith("text")]
    assert lines == ["x" * 24, "x" * 23 + "\u2026"]  # two lines, the second truncated


# -- CSV ----------------------------------------------------------------------


@pytest.mark.parametrize("value", ["=1+1", "+cmd", "-2", "@SUM(A1)", "\tx", "\rx", "  =padded"])
def test_safe_cell_neutralises_formula_starts(value):
    assert safe_cell(value).startswith("'")


@pytest.mark.parametrize("value", ["plain", "Web -> DB", "1 = 1", ""])
def test_safe_cell_leaves_ordinary_text_alone(value):
    assert safe_cell(value) == value


def test_threat_columns_are_stable():
    assert THREAT_COLUMNS[:5] == ("element", "element_type", "methodology", "category", "title")


# -- text ---------------------------------------------------------------------


def test_clean_text_strips_xml_illegal_characters_only():
    assert clean_text("a\x00b\x0bc\x1fd￾e") == "abcde"
    assert clean_text("tab\tnewline\nok é") == "tab\tnewline\nok é"
    assert clean_text(None) == ""


# -- data source labels never carry credentials ---------------------------------


@pytest.mark.parametrize(
    ("url", "expected"),
    [
        ("https://github.com/org/repo", "https://github.com/org/repo"),
        ("https://user:ghp_SECRET@github.com/org/repo/", "https://github.com/org/repo"),
        ("https://ghp_SECRET@github.com/org/repo?token=ghp_SECRET#frag", "https://github.com/org/repo"),
        ("not a url", ""),
        ("", ""),
    ],
)
def test_safe_repo_url_removes_credentials_query_and_fragment(url, expected):
    assert safe_repo_url(url) == expected


def _repo_with_run(tmp_path, details):
    repo = SQLRepository(f"sqlite:///{tmp_path}/r.db")
    now = datetime.now(UTC)
    repo.save_application(Application(id="APP-1", name="Shop", owner_user_id="U", created_at=now, updated_at=now))
    repo.save_data_source(DataSourceConfig(
        id="DS-1", application_id="APP-1", type="github_repository", connection_details=details,
        created_at=now,
    ))
    repo.save_data_source(DataSourceConfig(
        id="DS-2", application_id="APP-1", type="manifest_upload",
        connection_details={"filename": "deploy.yaml", "content": "SECRET-CONTENT"}, created_at=now,
    ))
    repo.save_run(ThreatModelRun(
        id="RUN-1", application_id="APP-1", triggered_by_user_id="U", data_source_config_ids=["DS-1", "DS-2"],
        status="completed", started_at=now, completed_at=now,
    ))
    return repo


def test_report_lists_sources_by_type_and_safe_label_never_credentials_or_content(tmp_path):
    repo = _repo_with_run(tmp_path, {
        "repo_url": "https://user:ghp_SECRET@github.com/org/repo", "github_token": "ghp_SECRET", "ref": "main",
    })

    report = build_report("APP-1", "RUN-1", repo)
    everything = render_html(report) + threats_csv(report) + report.model_dump_json()

    assert [(s.type, s.label) for s in report.sources] == [
        ("github_repository", "https://github.com/org/repo @ main"),
        ("manifest_upload", "deploy.yaml"),
    ]
    assert "ghp_SECRET" not in everything and "SECRET-CONTENT" not in everything
    assert "ghp_SECRET" not in render_docx(report).decode("latin-1")


def test_report_includes_model_metadata_in_html_and_docx(tmp_path):
    """Model metadata : owner,
    reviewer, contributors flow from Application through to both renderers."""
    repo = SQLRepository(f"sqlite:///{tmp_path}/r.db")
    now = datetime.now(UTC)
    repo.save_application(Application(
        id="APP-1", name="Shop", owner_user_id="U", created_at=now, updated_at=now,
        owner="Jane Owner", reviewer="Pat Reviewer", contributors=["Alex", HOSTILE],
    ))
    repo.save_run(ThreatModelRun(
        id="RUN-1", application_id="APP-1", triggered_by_user_id="U", status="completed",
        started_at=now, completed_at=now,
    ))

    report = build_report("APP-1", "RUN-1", repo)

    assert report.owner == "Jane Owner"
    assert report.reviewer == "Pat Reviewer"
    assert report.contributors == ["Alex", HOSTILE]
    html = render_html(report)
    assert "Jane Owner" in html and "Pat Reviewer" in html and "Alex" in html
    assert "<img src=x onerror=alert(1)>" not in html  # autoescaped, not stripped
    document = Document(io.BytesIO(render_docx(report)))
    docx_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Jane Owner" in docx_text and "Pat Reviewer" in docx_text


def test_build_report_rejects_unknown_application_and_foreign_run(tmp_path):
    from kresheen.engine.threats.service import UnknownRunError

    repo = _repo_with_run(tmp_path, {})
    with pytest.raises(LookupError):
        build_report("APP-NOPE", "RUN-1", repo)
    with pytest.raises(UnknownRunError):
        build_report("APP-1", "RUN-NOPE", repo)


# -- statements ---------------------------------------------------------------


def test_statements_use_the_spec_labels_and_skip_out_of_scope():
    lines = statements_for("data_store", {"is_encrypted": False, "stores_credentials": True, "out_of_scope": True,
                                          "out_of_scope_reason": "vendor"})

    assert lines == ["Encrypted at rest: No", "Stores credentials: Yes"]
    assert statements_for("data_flow", {"protocol": "gRPC"}) == ["Protocol: gRPC"]


# -- SBOM ---------------------------------------------------------------------


def test_sbom_rows_resolve_package_severity_and_order_worst_first():
    sbom = {
        "components": [{"bom-ref": "r1", "name": "lodash", "version": "4.17.20"}, {"bom-ref": "r2", "name": "left-pad"}],
        "vulnerabilities": [
            {"id": "CVE-2", "ratings": [{"score": 1}, {"severity": "low"}], "affects": [{"ref": "r2"}]},
            {"id": "CVE-1", "ratings": [{"severity": "critical"}], "affects": [{"ref": "r1"}]},
            "junk",
        ],
    }

    rows = sbom_vulnerability_rows(sbom)

    assert rows == [
        {"id": "CVE-1", "severity": "critical", "package": "lodash@4.17.20"},
        {"id": "CVE-2", "severity": "low", "package": "left-pad"},
    ]
    assert sbom_vulnerability_rows(None) == [] and len(sbom_vulnerability_rows(
        {"vulnerabilities": [{"id": str(i)} for i in range(80)]}, limit=50)) == 50


# -- DOCX / HTML ----------------------------------------------------------------


def test_docx_handles_control_characters_and_hostile_text(tmp_path):
    repo = _repo_with_run(tmp_path, {})
    report = build_report("APP-1", "RUN-1", repo)
    dirty = report.model_copy(update={"application_name": "Shop\x00\x0b " + HOSTILE})

    data = render_docx(dirty)

    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "Shop " in text and "<img src=x" in text and "\x00" not in text


# -- PDF ------------------------------------------------------------------------


def test_pdf_is_well_formed_and_never_carries_credentials(tmp_path):
    repo = _repo_with_run(tmp_path, {
        "repo_url": "https://user:ghp_SECRET@github.com/org/repo", "github_token": "ghp_SECRET", "ref": "main",
    })
    report = build_report("APP-1", "RUN-1", repo)

    data = render_pdf(report)

    assert data.startswith(b"%PDF-")
    assert b"ghp_SECRET" not in data


def test_pdf_handles_control_characters_and_hostile_text(tmp_path):
    repo = _repo_with_run(tmp_path, {})
    report = build_report("APP-1", "RUN-1", repo)
    dirty = report.model_copy(update={"application_name": "Shop\x00\x0b " + HOSTILE})

    data = render_pdf(dirty)  # raises if reportlab's Paragraph markup parser chokes on the escaped text

    assert data.startswith(b"%PDF-")


# -- SARIF ------------------------------------------------------------------------


def _finding(id_, severity, title="Vulnerable dependency", element="api"):
    from kresheen.reporting.model import ReportFinding

    return ReportFinding(
        id=id_, title=title, severity=severity, status="open", domain="vulnerability",
        source_tool="grype", element=element, description="A vulnerable package version is in use.",
        discovered_at=datetime.now(UTC),
    )


def test_build_sarif_maps_severity_to_level_and_never_leaks_control_chars(tmp_path):
    repo = _repo_with_run(tmp_path, {})
    report = build_report("APP-1", "RUN-1", repo)
    findings = [
        _finding("F-1", "critical", title="Bad\x00 " + HOSTILE),
        _finding("F-2", "medium"),
        _finding("F-3", "low"),
    ]
    report = report.model_copy(update={"findings": findings})

    sarif = build_sarif(report)

    assert sarif["version"] == "2.1.0"
    assert sarif["runs"][0]["tool"]["driver"]["name"] == "Kresheen"
    results = sarif["runs"][0]["results"]
    assert len(results) == 3
    assert [r["level"] for r in results] == ["error", "warning", "note"]
    assert all(r["ruleId"] == "grype" for r in results)
    assert "\x00" not in results[0]["properties"]["title"]
    assert "<img src=x" in results[0]["properties"]["title"]  # cleaned of control chars, not stripped of markup


# -- Threat Dragon JSON export ---------------------------------------------------


def test_build_threat_dragon_model_places_components_flows_and_boundaries(tmp_path):
    repo = _repo_with_run(tmp_path, {})
    application = repo.get_application("APP-1")
    components = [
        ReportComponent(id="C1", name="api", type="process", trust_boundary="prod", worst_severity=None,
                         open_finding_count=0, is_manual=False),
        ReportComponent(id="C2", name="db", type="data_store", trust_boundary="prod", worst_severity=None,
                         open_finding_count=0, is_manual=False),
        ReportComponent(id="C3", name=HOSTILE, type="external_entity", trust_boundary=None, worst_severity=None,
                         open_finding_count=0, is_manual=False),
    ]
    flows = [ReportFlow(id="F1", source_id="C3", target_id="C1", source_name="ext", target_name="api",
                         label="request", is_manual=False)]
    boundaries = [ReportBoundary(name="prod", position_x=10.0, position_y=20.0)]

    model = build_threat_dragon_model(application, components, flows, boundaries)

    cells = model["detail"]["diagrams"][0]["cells"]
    shapes_by_id = {c["id"]: c["shape"] for c in cells}
    assert shapes_by_id["C1"] == "process"
    assert shapes_by_id["C2"] == "store"
    assert shapes_by_id["C3"] == "actor"
    assert shapes_by_id["boundary:prod"] == "trust-boundary-box"
    flow_cell = next(c for c in cells if c["id"] == "F1")
    assert flow_cell["source"] == {"cell": "C3"} and flow_cell["target"] == {"cell": "C1"}
    actor_cell = next(c for c in cells if c["id"] == "C3")
    # JSON has no markup to escape into, so hostile text is passed through as-is
    # once XML-illegal control characters are stripped -- confirm the latter.
    assert actor_cell["data"]["name"] == HOSTILE
    assert "\x00" not in str(model)


def test_build_threat_dragon_model_via_report_endpoint_inputs(tmp_path):
    """Same inputs api_report.py's export_threat_dragon endpoint builds from
    (dfd_parts over build_run_dfd) round-trip through the exporter cleanly."""
    repo = _repo_with_run(tmp_path, {})
    application = repo.get_application("APP-1")
    from kresheen.engine.dfd.diagrams import get_or_create_default_diagram
    from kresheen.engine.dfd.view import build_run_dfd

    diagram_id = get_or_create_default_diagram("APP-1", repo).id
    nodes, edges, boundaries = build_run_dfd("APP-1", "RUN-1", repo, diagram_id)
    components, flows, report_boundaries = dfd_parts(nodes, edges, boundaries)

    model = build_threat_dragon_model(application, components, flows, report_boundaries)

    assert model["version"] == "2.4.0"
    assert model["summary"]["title"] == "Shop"
    assert isinstance(model["detail"]["diagrams"][0]["cells"], list)


def test_csv_round_trips_multiline_and_quotes(tmp_path):
    repo = _repo_with_run(tmp_path, {})
    report = build_report("APP-1", "RUN-1", repo)

    rows = list(csv.reader(io.StringIO(threats_csv(report).lstrip("\ufeff"))))

    assert tuple(rows[0]) == THREAT_COLUMNS
    assert all(len(row) == len(THREAT_COLUMNS) for row in rows)
